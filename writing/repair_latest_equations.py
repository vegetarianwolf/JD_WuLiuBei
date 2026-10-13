"""Minimal, assertion-guarded repairs of confirmed OMML nesting errors.

Usage (the caller decides whether and where to save):
    from docx import Document
    doc = Document(path)
    report = fix_equation_30(doc)

This module never loads or saves a document.  fix_equation_30 only changes the
numbered equation (30), in memory.  It preserves every existing XML node, text,
font/style property, paragraph property, tab, and equation number.  Two existing
element tags are renamed, seven existing equation-level terms are reparented,
and two mathematical symbols (+ and l) with copied formatting are added.

Intended formula, established by the surrounding node/vehicle definitions:
    r_{ih}^{kl} = 1 => T_{eta_{ih}^{+}}^{l} >= T_{eta_{ih}^{-}}^{k},
    i in R, h in V_i^r, k,l in D.

The independent fix_equation_28_context entry point repairs only equation (28)
and its preceding outbound-node sequence, preserving all mathematical symbols.
"""

from copy import deepcopy
import re

from lxml import etree


M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def _m(local):
    return f"{{{M}}}{local}"


def _math_text(element):
    if element is None:
        return ""
    return "".join(node.text or "" for node in element.iter(_m("t")))


def _require(condition, message):
    # Deliberately not a Python assert: safeguards must also run under python -O.
    if not condition:
        raise AssertionError(f"Known OMML repair refused: {message}")


def fix_equation_30(doc):
    """Repair only the known malformed equation (30), returning an audit report.

    All source-shape assertions run before any mutation.  An already corrected
    or otherwise different equation is rejected rather than silently rewritten.
    The caller should save to a separate file and visually render the result.
    """
    candidates = [
        (index, paragraph)
        for index, paragraph in enumerate(doc.paragraphs)
        if re.fullmatch(r"\s*\(30\)\s*", paragraph.text)
    ]
    _require(len(candidates) == 1, "expected exactly one numbered (30) paragraph")
    paragraph_index, paragraph = candidates[0]
    p = paragraph._p
    math_nodes = p.findall(_m("oMath"))
    _require(len(math_nodes) == 1, "expected one direct oMath child")
    math = math_nodes[0]
    terms = list(math)
    _require(
        [node.tag for node in terms]
        == [_m("sSubSup"), _m("r"), _m("sSub")],
        "expected relay variable, implication run, malformed T subscript",
    )
    relay, implication, malformed_t = terms
    _require(_math_text(relay) == "rihkl", "unexpected relay variable")
    _require(_math_text(implication) == "=1⇒", "unexpected implication")
    _require(
        [node.tag for node in malformed_t]
        == [_m("sSubPr"), _m("e"), _m("sub")],
        "unexpected malformed T structure",
    )
    properties, base, subscript = list(malformed_t)
    _require(_math_text(base) == "T", "unexpected first time-variable base")
    subterms = list(subscript)
    _require(
        [node.tag for node in subterms]
        == [
            _m("sSubSup"), _m("r"), _m("sSubSup"), _m("r"),
            _m("sSubSup"), _m("r"), _m("r"), _m("r"), _m("ctrlPr"),
        ],
        "unexpected contents swallowed by the first T subscript",
    )
    eta, relation, rhs_t, domain_start, domain_v, domain_k, domain_l, domain_end, _ = subterms
    _require(_math_text(eta) == "ηih", "unexpected pickup-node expression")
    _require(_math_text(relation) == "≥", "unexpected comparison operator")
    _require(_math_text(rhs_t) == "Tηih−k", "unexpected unloading time")
    _require(
        [_math_text(node) for node in (domain_start, domain_v, domain_k, domain_l, domain_end)]
        == [",i∈R,h∈", "Vir", ",k", ",l∈", "D."],
        "unexpected variable-domain declaration",
    )
    eta_sup = eta.find(_m("sup"))
    _require(eta_sup is not None, "pickup node has no superscript container")
    _require(
        [node.tag for node in eta_sup] == [_m("ctrlPr")]
        and _math_text(eta_sup) == "",
        "pickup superscript is not the known empty placeholder",
    )
    rhs_eta = rhs_t.find(f'{_m("sub")}/{_m("sSubSup")}')
    _require(rhs_eta is not None, "unloading time has no eta-node subscript")
    minus_runs = rhs_eta.findall(f'{_m("sup")}/{_m("r")}')
    _require(
        len(minus_runs) == 1 and _math_text(minus_runs[0]) == "−",
        "expected one formatted minus run to copy",
    )
    l_runs = [
        node for node in relay.findall(f'{_m("sup")}/{_m("r")}')
        if _math_text(node) == "l"
    ]
    _require(len(l_runs) == 1, "expected one formatted l run to copy")
    relay_sup = relay.find(_m("sup"))
    relay_sup_control = relay_sup.find(_m("ctrlPr"))
    _require(relay_sup_control is not None, "relay superscript lacks formatting control")

    # Prepare additions before mutation, keeping the existing font/script style.
    plus_run = deepcopy(minus_runs[0])
    plus_text = plus_run.find(_m("t"))
    _require(plus_text is not None, "minus run lacks a math text element")
    plus_text.text = "+"
    vehicle_sup = etree.Element(_m("sup"))
    vehicle_sup.append(deepcopy(l_runs[0]))
    vehicle_sup.append(deepcopy(relay_sup_control))

    original_nodes = tuple(p.iter())
    original_texts = {node: (node.text, node.tail) for node in original_nodes}
    original_attributes = {node: dict(node.attrib) for node in original_nodes}
    original_tags = {node: node.tag for node in original_nodes}
    original_number = paragraph.text

    # Only the eta node remains under the first T.  Preserve its original ctrlPr.
    promoted_terms = subterms[1:-1]
    for node in promoted_terms:
        math.append(node)
    malformed_t.tag = _m("sSubSup")
    properties.tag = _m("sSubSupPr")
    malformed_t.append(vehicle_sup)
    eta_sup.insert(0, plus_run)

    # Verify preservation plus the exact, limited mathematical correction.
    final_nodes = set(p.iter())
    _require(all(node in final_nodes for node in original_nodes), "an original XML node was lost")
    _require(
        all((node.text, node.tail) == original_texts[node] for node in original_nodes),
        "an existing text value changed",
    )
    _require(
        all(dict(node.attrib) == original_attributes[node] for node in original_nodes),
        "an existing attribute changed",
    )
    allowed_renames = {malformed_t: _m("sSubSup"), properties: _m("sSubSupPr")}
    _require(
        all(node.tag == allowed_renames.get(node, original_tags[node]) for node in original_nodes),
        "an unrelated element tag changed",
    )
    _require(paragraph.text == original_number, "paragraph tabs or equation number changed")
    expected = "rihkl=1⇒Tηih+l≥Tηih−k,i∈R,h∈Vir,k,l∈D."
    _require(_math_text(math) == expected, "repaired token sequence differs from the intended formula")

    return {
        "equation": 30,
        "paragraph_index": paragraph_index,
        "paragraph_id": p.get(f"{{{W14}}}paraId"),
        "retained_original_xml_nodes": len(original_nodes),
        "promoted_original_terms": len(promoted_terms),
        "renamed_elements": ["sSub -> sSubSup", "sSubPr -> sSubSupPr"],
        "added_math_symbols": ["+", "l"],
        "result_math_tokens": expected,
        "saved": False,
    }


def fix_equation_28_context(doc):
    """Repair only the two confirmed nesting errors around equation (28).

    The preceding sentence's outbound sequence is eta_{ih}^{+}, i+N.
    Its malformed superscript run '+,i+N' is split into an existing '+' run
    and a copied-format ',i+N' run at equation level.  Equation (28)'s pickup
    load becomes q_{eta_{ih}^{+}}=+1,q_{i+N}=-1 by promoting three existing
    terms out of its incorrect subscript container.

    No mathematical symbol is inserted or removed.  Every original XML node
    remains; the sole existing text-node change is the explicitly checked run
    split.  Both complete input shapes are checked before either mutation.
    No paragraph property, ordinary text, tab, or equation number is changed.
    """
    candidates = [
        (index, paragraph)
        for index, paragraph in enumerate(doc.paragraphs)
        if re.fullmatch(r"\s*\(28\)\s*", paragraph.text)
    ]
    _require(len(candidates) == 1, "expected exactly one numbered (28) paragraph")
    paragraph_index, equation_paragraph = candidates[0]
    _require(paragraph_index > 0, "equation (28) has no preceding sentence")
    context_paragraph = doc.paragraphs[paragraph_index - 1]
    _require(
        context_paragraph.text
        == "对于接力送任务，在中继站设置卸货节点和取货节点。入站段依次访问，出站段依次访问，相应载荷增量为：",
        "unexpected sentence immediately preceding equation (28)",
    )
    context_math = context_paragraph._p.findall(_m("oMath"))
    _require(
        [_math_text(node) for node in context_math]
        == ["h", "ηih−", "ηih+", "i,ηih−", "ηih+,i+N"],
        "unexpected inline formulas preceding equation (28)",
    )
    outbound = context_math[-1]
    _require(
        [node.tag for node in outbound] == [_m("sSubSup")],
        "outbound sequence is not the known malformed eta expression",
    )
    outbound_eta = outbound[0]
    _require(
        _math_text(outbound_eta.find(_m("e"))) == "η"
        and _math_text(outbound_eta.find(_m("sub"))) == "ih",
        "unexpected outbound pickup node",
    )
    outbound_sup = outbound_eta.find(_m("sup"))
    _require(
        outbound_sup is not None
        and [node.tag for node in outbound_sup] == [_m("r"), _m("ctrlPr")],
        "unexpected outbound superscript structure",
    )
    compound_run = outbound_sup[0]
    compound_texts = compound_run.findall(_m("t"))
    _require(
        len(compound_texts) == 1 and compound_texts[0].text == "+,i+N",
        "outbound superscript is not the known compound run",
    )
    compound_text = compound_texts[0]

    equation_nodes = equation_paragraph._p.findall(_m("oMath"))
    _require(len(equation_nodes) == 1, "expected one direct oMath in equation (28)")
    equation_math = equation_nodes[0]
    _require(
        [node.tag for node in equation_math]
        == [_m("sSub"), _m("r"), _m("sSub"), _m("r"), _m("sSub")],
        "equation (28) lacks the known nested pickup-load expression",
    )
    _require(
        [_math_text(node) for node in equation_math]
        == ["qi", "=+1,", "qηih−", "=−1,", "qηih+=+1,qi+N=−1."],
        "unexpected load values or nodes in equation (28)",
    )
    pickup_q = equation_math[-1]
    _require(_math_text(pickup_q.find(_m("e"))) == "q", "unexpected pickup-load base")
    pickup_sub = pickup_q.find(_m("sub"))
    _require(
        pickup_sub is not None
        and [node.tag for node in pickup_sub]
        == [_m("sSubSup"), _m("r"), _m("sSub"), _m("r"), _m("ctrlPr")],
        "unexpected terms inside the pickup-load subscript",
    )
    pickup_terms = list(pickup_sub)
    _require(
        [_math_text(node) for node in pickup_terms]
        == ["ηih+", "=+1,", "qi+N", "=−1.", ""],
        "unexpected swallowed load terms",
    )
    _require(
        _math_text(pickup_terms[2].find(_m("e"))) == "q"
        and _math_text(pickup_terms[2].find(_m("sub"))) == "i+N",
        "unexpected final-delivery load variable",
    )

    # Prepare the copied-format suffix before changing either paragraph.
    suffix_run = deepcopy(compound_run)
    suffix_run.find(_m("t")).text = ",i+N"
    paragraphs = [context_paragraph, equation_paragraph]
    original_nodes = tuple(node for paragraph in paragraphs for node in paragraph._p.iter())
    original_texts = {node: (node.text, node.tail) for node in original_nodes}
    original_attributes = {node: dict(node.attrib) for node in original_nodes}
    original_tags = {node: node.tag for node in original_nodes}
    original_plain_texts = [paragraph.text for paragraph in paragraphs]
    original_math_tokens = [
        [_math_text(node) for node in paragraph._p.findall(_m("oMath"))]
        for paragraph in paragraphs
    ]

    compound_text.text = "+"
    outbound.append(suffix_run)
    promoted_terms = pickup_terms[1:-1]
    for node in promoted_terms:
        equation_math.append(node)

    final_nodes = {node for paragraph in paragraphs for node in paragraph._p.iter()}
    _require(all(node in final_nodes for node in original_nodes), "a context/equation (28) node was lost")
    _require(
        all(
            (node.text, node.tail)
            == (("+", original_texts[node][1]) if node is compound_text else original_texts[node])
            for node in original_nodes
        ),
        "text changed outside the expressly permitted superscript-run split",
    )
    _require(
        all(node.tag == original_tags[node] for node in original_nodes)
        and all(dict(node.attrib) == original_attributes[node] for node in original_nodes),
        "an original context/equation (28) tag or attribute changed",
    )
    _require(
        [paragraph.text for paragraph in paragraphs] == original_plain_texts,
        "ordinary context text, paragraph tabs, or equation (28) number changed",
    )
    _require(
        [
            [_math_text(node) for node in paragraph._p.findall(_m("oMath"))]
            for paragraph in paragraphs
        ] == original_math_tokens,
        "a mathematical symbol was added, removed, or reordered",
    )

    return {
        "equation": 28,
        "paragraph_indices": [paragraph_index - 1, paragraph_index],
        "paragraph_ids": [paragraph._p.get(f"{{{W14}}}paraId") for paragraph in paragraphs],
        "retained_original_xml_nodes": len(original_nodes),
        "split_math_run": {"before": "+,i+N", "superscript": "+", "promoted_suffix": ",i+N"},
        "promoted_original_terms": len(promoted_terms),
        "added_math_symbols": [],
        "removed_math_symbols": [],
        "saved": False,
    }
